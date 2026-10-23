# -*- coding: utf-8 -*-
import json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

D = '/Users/aev/WorkBuddy/WYDoc/01-功能盘点'
rows = json.load(open(D + '/data/mapped.json'))

# 排序：域 A-E，一级模块顺序
DOM_ORDER = ['A 资产', 'B 费用', 'C 办公', 'D 物联', 'E 平台底座']
L1_ORDER = [
    'A1 空间资源', 'A2 人员档案', 'A3 车辆与车位', 'A4 生命周期', 'A5 经营合同', 'A6 通行与出入',
    'B1 收费配置', 'B2 计费与出账', 'B3 抄表与公摊', 'B4 收款与票据', 'B5 退费与票据', 'B6 欠费与催缴',
    'C1 审批流', 'C2 考勤', 'C3 工单与服务', 'C4 现场作业', 'C5 采购与库存', 'C6 行政与社区',
    'D1 设备与台账', 'D2 门禁', 'D3 车场设备', 'D4 视频监控', 'D7 边缘设备',
    'E1 组织与权限', 'E2 系统配置', 'E3 报表与驾驶舱', 'E4 端与渠道配置',
]
rows.sort(key=lambda r: (DOM_ORDER.index(r['domain']), L1_ORDER.index(r['l1']) if r['l1'] in L1_ORDER else 99))

wb = Workbook()
ws = wb.active
ws.title = '功能盘点与研判'

NAVY = '1F3864'
HDR_FILL = PatternFill('solid', fgColor=NAVY)
HDR_FONT = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name='微软雅黑', size=9)
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DOM_COLOR = {
    'A 资产': 'E8F0FE', 'B 费用': 'FDF2E9', 'C 办公': 'EAF7EE',
    'D 物联': 'F3E8FD', 'E 平台底座': 'F2F2F2',
}
ACT_COLOR = {'保留': 'C6EFCE', '升级': 'FFEB9C', '合并': 'DDEBF7', '重建': 'F8CBAD', '放弃': 'E7E6E6'}

headers = ['序号', '顶级域', '一级模块', '二级功能点', '现状归属', '现状功能点',
           '现状页面路径', '处置建议', '目标端', '优先级', '升级 / 重构要点']
widths = [6, 12, 18, 16, 10, 14, 46, 10, 8, 8, 72]

for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30

r = 2
for idx, x in enumerate(rows, start=1):
    vals = [idx, x['domain'], x['l1'], x['name'], x['group'] + '/' + x['module'],
            x['name'], x['href'] or '—', x['act'], x['target'], x['prio'], x['note']]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = BODY_FONT; c.border = BORDER
        c.alignment = Alignment(vertical='center', wrap_text=(ci in (7, 11)),
                                horizontal='center' if ci in (1, 8, 9, 10) else 'left')
    ws.cell(row=r, column=2).fill = PatternFill('solid', fgColor=DOM_COLOR[x['domain']])
    ws.cell(row=r, column=8).fill = PatternFill('solid', fgColor=ACT_COLOR[x['act']])
    ws.row_dimensions[r].height = max(30, math.ceil(len(x['note']) / 34) * 15)
    r += 1

ws.freeze_panes = 'D2'
ws.auto_filter.ref = 'A1:K%d' % (r - 1)

# ============ Sheet2 汇总 ============
ws2 = wb.create_sheet('域汇总')
idx_map = {}
for x in rows:
    k = (x['domain'], x['l1'])
    idx_map.setdefault(k, {'n': 0, '保留': 0, '升级': 0, '合并': 0, '重建': 0, '放弃': 0, 'P0': 0, 'P1': 0, 'P2': 0})
    d = idx_map[k]
    d['n'] += 1; d[x['act']] += 1; d[x['prio']] += 1

h2 = ['顶级域', '一级模块', '功能点数', '保留', '升级', '合并', '重建', '放弃', 'P0', 'P1', 'P2']
w2 = [14, 20, 10, 8, 8, 8, 8, 8, 8, 8, 8]
for i, (h, w) in enumerate(zip(h2, w2), start=1):
    c = ws2.cell(row=1, column=i, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws2.column_dimensions[get_column_letter(i)].width = w

rr = 2
for dom in DOM_ORDER:
    sub = {k: v for k, v in idx_map.items() if k[0] == dom}
    for (dm, l1), v in sorted(sub.items(), key=lambda kv: L1_ORDER.index(kv[0][1]) if kv[0][1] in L1_ORDER else 99):
        vals = [dm, l1, v['n'], v['保留'], v['升级'], v['合并'], v['重建'], v['放弃'], v['P0'], v['P1'], v['P2']]
        for ci, val in enumerate(vals, start=1):
            c = ws2.cell(row=rr, column=ci, value=val)
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal='center' if ci != 2 else 'left', vertical='center')
        ws2.cell(row=rr, column=1).fill = PatternFill('solid', fgColor=DOM_COLOR[dom])
        rr += 1
    # 小计
    tot = [v for (dm, l1), v in sub.items()]
    vals = [dom, '小计', sum(v['n'] for v in tot), sum(v['保留'] for v in tot), sum(v['升级'] for v in tot),
            sum(v['合并'] for v in tot), sum(v['重建'] for v in tot), sum(v['放弃'] for v in tot),
            sum(v['P0'] for v in tot), sum(v['P1'] for v in tot), sum(v['P2'] for v in tot)]
    for ci, val in enumerate(vals, start=1):
        c = ws2.cell(row=rr, column=ci, value=val)
        c.font = Font(name='微软雅黑', size=9, bold=True); c.border = BORDER
        c.fill = PatternFill('solid', fgColor='D9E1F2')
        c.alignment = Alignment(horizontal='center' if ci != 2 else 'left', vertical='center')
    rr += 1

# ============ Sheet3 新增建议 ============
NEW = [
    ('A 资产', '空间与资产主数据平台', '统一房屋/车位/商铺/仓库/场地为"空间单元"，一物一码', 'P0'),
    ('A 资产', '客户关系档案（CRM）', '业主/家属/租户/商户统一客户视图，一房多人、一人多房', 'P0'),
    ('A 资产', '资产全生命周期', '交付-装修-使用-维修-改造-报废，含资产卡与折旧', 'P2'),
    ('A 资产', '多经资源经营', '广告位、场地、公共区域经营与合同收益管理', 'P2'),
    ('B 费用', '计费规则引擎', '费项+规则+范围可配置，支持阶梯/公摊/滞纳金/免租期', 'P0'),
    ('B 费用', '统一收款台', '线上支付、线下收款、代扣、预存抵扣、退款一站式', 'P0'),
    ('B 费用', '催缴策略引擎', '按账龄自动分级催缴（短信/电话/上门/停服），输出催缴率', 'P1'),
    ('B 费用', '业财一体', '账单-收款-凭证自动生成，对接金蝶/用友/税控', 'P1'),
    ('B 费用', '电子发票', '数电票自动开具、红冲、合并开票', 'P1'),
    ('B 费用', '专项维修资金', '大修基金归集、支用审批、公示', 'P2'),
    ('C 办公', '统一工单中心', '报修/投诉/巡检异常/设备告警/内部任务同一工单模型 + SLA', 'P0'),
    ('C 办公', '统一审批引擎', 'BPMN 流程引擎 + 表单设计器，覆盖采购/调拨/合同/装修/减免', 'P0'),
    ('C 办公', '移动办公端', '巡检打点、工单处理、抄表、打卡、审批一体化 App', 'P0'),
    ('C 办公', '智能派单', '按技能、位置、负载、历史质量自动派单', 'P2'),
    ('C 办公', '服务质量体系', 'SLA、满意度 NPS、员工绩效与扣罚联动', 'P1'),
    ('D 物联', 'IoT 接入平台', '设备影子、物模型、MQTT/HTTP/GB28181 协议适配、离线缓存', 'P0'),
    ('D 物联', '统一告警中心', '设备/能耗/安全/门禁异常统一告警 + 自动转工单', 'P0'),
    ('D 物联', '门禁与访客', '人脸/二维码/蓝牙开门，访客邀约 - 授权 - 核销闭环', 'P1'),
    ('D 物联', '车场与车牌识别', '道闸、车牌识别、无感支付、车位共享、访客停车', 'P1'),
    ('D 物联', '能耗管理', '远传表计、分项计量、能耗基线、异常用能告警', 'P1'),
    ('E 平台底座', '多项目集团管控', '集团-区域-项目三级组织与数据隔离、指标横向对比', 'P0'),
    ('E 平台底座', '权限矩阵', 'RBAC + 数据权限 + 字段权限，可视化配置与冲突检测', 'P0'),
    ('E 平台底座', '报表中心与驾驶舱', '维度+指标自助取数，看板可配置', 'P1'),
    ('E 平台底座', '开放平台/API 网关', '对接第三方（支付/税控/门禁/视频/政府平台）', 'P1'),
    ('E 平台底座', '主数据与字典管理', '统一编码体系，避免各模块自建字典', 'P1'),
    ('E 平台底座', '审计与合规', '操作留痕、登录风控、数据导出审批', 'P1'),
    ('E 平台底座', 'AI 助手', '工单摘要、智能客服、收费预测、报告生成', 'P2'),
]
ws3 = wb.create_sheet('新增建议')
h3 = ['序号', '所属域', '建议新增模块', '解决什么问题', '优先级']
w3 = [6, 14, 26, 62, 8]
for i, (h, w) in enumerate(zip(h3, w3), start=1):
    c = ws3.cell(row=1, column=i, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws3.column_dimensions[get_column_letter(i)].width = w
for i, (dm, nm, why, p) in enumerate(NEW, start=1):
    vals = [i, dm, nm, why, p]
    for ci, v in enumerate(vals, start=1):
        c = ws3.cell(row=i + 1, column=ci, value=v)
        c.font = BODY_FONT; c.border = BORDER
        c.alignment = Alignment(horizontal='center' if ci in (1, 5) else 'left', vertical='center', wrap_text=(ci == 4))
    ws3.cell(row=i + 1, column=2).fill = PatternFill('solid', fgColor=DOM_COLOR[dm])
    ws3.row_dimensions[i + 1].height = 26
ws3.freeze_panes = 'A2'

# ============ Sheet4 说明 ============
ws4 = wb.create_sheet('口径与结论')
ws4.column_dimensions['A'].width = 22
ws4.column_dimensions['B'].width = 100

def put(r, k, v, bold=False, fill=None):
    a = ws4.cell(row=r, column=1, value=k)
    a.font = Font(name='微软雅黑', size=10, bold=True)
    a.alignment = Alignment(vertical='top')
    if fill: a.fill = PatternFill('solid', fgColor=fill)
    b = ws4.cell(row=r, column=2, value=v)
    b.font = Font(name='微软雅黑', size=10, bold=bold)
    b.alignment = Alignment(vertical='top', wrap_text=True)
    ws4.row_dimensions[r].height = max(20, math.ceil(len(str(v)) / 52) * 16)

from collections import Counter
cact = Counter([x['act'] for x in rows])
cdom = Counter([x['domain'] for x in rows])

put(1, '数据来源', 'HC 智慧物业 PC 管理后台（demo.homecommunity.cn）真实登录抓取：菜单接口 menu.listCatalog / menu.listCatalogMenus，共 13 个一级菜单 + 设置下 2 个分组，合计 161 个二级菜单功能点；另参考前端路由表 432 个页面（含详情/编辑/打印页）作为补充。', True, 'D9E1F2')
put(2, '抓取时间', '2026-10-21')
put(3, '功能点口径', '以"后台侧边栏菜单项"为一个功能点；菜单内的详情页、编辑页、打印页不单独计数，但在"升级要点"中体现。')
put(4, '顶级分类口径', '按新系统规划：A 资产（人·车·房）/ B 费用 / C 办公（OA+现场作业）/ D 物联 / E 平台底座（跨域共享，不占用四大业务类但必须单列）。')
put(5, '处置口径', '保留=逻辑可用、换框架重做即可；升级=保留业务但需改交互或增强能力；合并=与其他功能重叠，收敛为一个模块；重建=业务逻辑或数据模型不成立，需要重新设计；放弃=价值低或风险高，建议不做。')
put(6, '优先级口径', 'P0=第一期必须完成（否则业务跑不起来）；P1=第二期；P2=可延后或按项目定制。')

put(8, '总量', '161 个功能点：合并 %d 项（%.0f%%）、保留 %d 项、升级 %d 项、重建 %d 项、放弃 %d 项。' % (
    cact['合并'], cact['合并'] / 161 * 100, cact['保留'], cact['升级'], cact['重建'], cact['放弃']), True, 'FFF2CC')
put(9, '域分布', '、'.join('%s %d 项' % (d, cdom[d]) for d in DOM_ORDER))
put(10, '核心结论 1', '菜单碎片化严重：一半功能应合并。典型如"费用"下 4 个出账入口（房屋/车位/商铺/合同）、"报表"下 11 张独立报表、"OA"下 5 个待办入口，本质是同一件事的多次实现，直接搬到新系统会把技术债一起搬过去。', True)
put(11, '核心结论 2', '资产域建模是最大风险：现有系统以"房屋"为中心挂业主，无法支撑一人多房、租户、商户、车位产权等场景。A1/A2 必须重建数据模型，否则后续费用、工单、门禁都会扭曲。', True)
put(12, '核心结论 3', '物联域最薄弱：原系统仅 11 个功能点，且分散在"设备/停车/小区"三个模块里，没有设备影子、没有告警中心、没有协议适配层。新系统物联要做的是"接入平台"而不是"设备列表"。', True)
put(13, '核心结论 4', 'OA 流程引擎必须换：现有自研流程（流程实例/我的流程/工作办理/工作单）无法支持会签、条件分支、加签、撤回，也无法被采购/合同/装修复用，建议直接采用成熟 BPMN 引擎 + 表单设计器。', True)
put(14, '核心结论 5', '办公域占比最高（62 项 / 38%%），说明原系统重心在"人怎么干活"，而不是"资产怎么增值"。重构时要区分：哪些是物业内部管理（考勤、采购、巡检），哪些是业主服务（报修、投诉、访客），后者应大幅前移到移动端。', True)

put(16, '下一步建议', '1）先对 P0 的 22 项逐条过会定版；2）A1/A2 出详细数据模型（实体+关系+状态机）；3）确定 BPMN 引擎与 IoT 平台选型；4）再按四大域出原型。')

wb.save(D + '/智慧物业PC后台_功能盘点与研判表.xlsx')
print('saved xlsx')
